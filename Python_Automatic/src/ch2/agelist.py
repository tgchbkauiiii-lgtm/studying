import openpyxl as ex
import datetime

book = ex.Workbook()
sheet = book.active

this_year = datetime.date.today().year

for i in range(80):
    age = i
    year = this_year - i
    age_cell = sheet.cell(i+1,1)
    age_cell.value = (f'{i}才')
    year_cell = sheet.cell(i+1,2)
    year_cell.value = (f'{year}年')

book.save("agelist.xlsx")