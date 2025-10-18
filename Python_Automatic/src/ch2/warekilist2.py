import openpyxl as xl

book = xl.Workbook()
sheet = book.active

sheet["a1"] = "西暦"
sheet["b1"] = "和暦"

start_y = 1845
for i in range(200):
    sei = start_y + i
    wa = f'=TEXT("{sei}/1/1","ggge年")'
    sheet.cell(row=(2+i),column=1,value=f'{sei}年')
    sheet.cell(row=(2+i),column=2,value=wa)
    print(sei,"=",wa)

book.save("wareki2.xlsx")
