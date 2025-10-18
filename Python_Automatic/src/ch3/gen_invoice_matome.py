import openpyxl as x,json

infile = "matome.json"
out_dir = "../../invoice02"
template_file = "../../invoice-template.xlsx"
subject = "2月分のご請求"

def gen_invoice():
    with open(infile,"rt") as fp:
        users = json.load(fp)
    for name, data in users.items():
        make_user_invoice(name, data)

def make_user_invoice(name, data):
    book = x.load_workbook(template_file)
    sheet = book.active
    sheet["b4"] = name
    sheet["c10"] = subject
    sheet["c11"] = data["total"]
    for i, it in enumerate(data["items"]):
        date, summary, cnt, price = it
        row = 15 + i
        sheet.cell(row,2,f'{summary}({date})')
        sheet.cell(row,5,cnt)
        sheet.cell(row,7,cnt*price)
    out_file = f'{out_dir}/{name}様.xlsx'
    book.save(out_file)
    print("save:",out_file)
    
if __name__ == "__main__":
    gen_invoice()